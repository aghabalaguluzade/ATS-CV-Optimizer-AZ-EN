import os
import re


def _add_list_paragraph(document, title, items):
    if not items:
        return
    document.add_heading(title, level=2)
    source_items = items if isinstance(items, list) else [items]
    normalized_items = []
    for item in source_items:
        if isinstance(item, str):
            normalized_items.extend(part.strip() for part in re.split(r"[|\n]+", item) if part.strip())
        elif item:
            normalized_items.append(str(item))
    for item in normalized_items:
        document.add_paragraph(item, style="List Bullet")


def generate_docx(cv_data, output_docx_path, language="en"):
    try:
        from docx import Document
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
    except Exception:
        print("[!] DOCX olusturma icin 'python-docx' gerekli. Kurulum: pip install python-docx")
        return False

    try:
        os.makedirs(os.path.dirname(output_docx_path), exist_ok=True)
        document = Document()
        for style in document.styles:
            if not hasattr(style, "font"):
                continue
            style.font.name = "Arial"
            r_pr = style._element.get_or_add_rPr()
            fonts = r_pr.rFonts
            if fonts is None:
                fonts = OxmlElement("w:rFonts")
                r_pr.insert(0, fonts)
            for font_type in ("ascii", "hAnsi", "eastAsia", "cs"):
                fonts.set(qn(f"w:{font_type}"), "Arial")
        labels = {
            "az": {
                "summary": "Peşəkar xülasə", "experience": "İş təcrübəsi",
                "projects": "Layihələr", "education": "Təhsil",
                "skills": "Texniki bacarıqlar", "certifications": "Sertifikatlar",
            },
            "en": {
                "summary": "Summary", "experience": "Work Experience",
                "projects": "Projects", "education": "Education",
                "skills": "Technical Skills", "certifications": "Certifications",
            },
        }.get(language, {
            "summary": "Summary", "experience": "Work Experience",
            "projects": "Projects", "education": "Education",
            "skills": "Technical Skills", "certifications": "Certifications",
        })

        document.add_heading(cv_data.get("name", ""), level=1)
        title = cv_data.get("title", "")
        if title:
            document.add_paragraph(title)

        contact_parts = [
            cv_data.get("location", ""),
            cv_data.get("email", ""),
            cv_data.get("phone", ""),
            cv_data.get("linkedin", ""),
            cv_data.get("github", ""),
        ]
        contact_text = " | ".join(part for part in contact_parts if part)
        if contact_text:
            document.add_paragraph(contact_text)

        summary = cv_data.get("summary", "")
        if summary:
            document.add_heading(labels["summary"], level=2)
            document.add_paragraph(summary)

        experience = cv_data.get("experience", [])
        if experience:
            document.add_heading(labels["experience"], level=2)
            for job in experience:
                role = job.get("role", "")
                company = job.get("company", "")
                period = job.get("period", "")
                location = job.get("location", "")
                header = " - ".join(part for part in [role, company] if part)
                if header:
                    document.add_paragraph(header)
                meta = " | ".join(part for part in [period, location] if part)
                if meta:
                    document.add_paragraph(meta)
                description = job.get("description", "")
                if description:
                    document.add_paragraph(description, style="List Bullet")

        projects = cv_data.get("projects", [])
        if projects:
            document.add_heading(labels["projects"], level=2)
            for project in projects:
                project_name = project.get("name", "")
                project_url = project.get("url", "")
                project_header = " - ".join(part for part in [project_name, project_url] if part)
                if project_header:
                    document.add_paragraph(project_header)
                project_description = project.get("description", "")
                if project_description:
                    document.add_paragraph(project_description, style="List Bullet")

        education = cv_data.get("education", [])
        if education:
            document.add_heading(labels["education"], level=2)
            for edu in education:
                school = edu.get("school", "")
                degree = edu.get("degree", "")
                period = edu.get("period", "")
                line = " | ".join(part for part in [school, degree, period] if part)
                if line:
                    document.add_paragraph(line)

        _add_list_paragraph(document, labels["skills"], cv_data.get("skills", []))
        _add_list_paragraph(document, labels["certifications"], cv_data.get("certifications", []))

        document.save(output_docx_path)
        return True
    except Exception as e:
        print(f"[!] DOCX olusturulurken hata: {e}")
        return False
